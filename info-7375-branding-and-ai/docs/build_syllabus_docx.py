from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

BASE = Path(__file__).parent
RED = 'C8102E'
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(.7)
s.left_margin = s.right_margin = Inches(.75)
s.footer_distance = Inches(.3)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Number']:
    st = doc.styles[name]
    st.font.name = 'Arial'
    st.font.color.rgb = RGBColor.from_string('000000')
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.05
doc.styles['Title'].font.size = Pt(21)
doc.styles['Title'].font.bold = True
doc.styles['Subtitle'].font.size = Pt(11)
doc.styles['Subtitle'].font.bold = True
doc.styles['Subtitle'].font.italic = False
for name, size in [('Heading 1',15), ('Heading 2',12), ('Heading 3',11)]:
    st = doc.styles[name]
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(RED)
    st.paragraph_format.space_before = Pt(13)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True

def inline(p, text):
    pattern = r'(\[[^\]]+\]\([^)]+\)|\*\*.+?\*\*|`[^`]+`)'
    for token in re.split(pattern, text):
        if not token:
            continue
        m = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', token)
        if m:
            h = OxmlElement('w:hyperlink')
            h.set(qn('r:id'), p.part.relate_to(m[2], RT.HYPERLINK, is_external=True))
            r = OxmlElement('w:r'); pr = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'), RED); pr.append(color)
            u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); pr.append(u)
            r.append(pr); t = OxmlElement('w:t'); t.text = m[1]; r.append(t); h.append(r); p._p.append(h)
        else:
            bold = token.startswith('**')
            text = token[2:-2] if bold else token.strip('`')
            run = p.add_run(text.replace('\\*', '*'))
            if bold: run.bold = True
    return p

def para(text, style=None):
    p = inline(doc.add_paragraph(style=style), text)
    if text.endswith(':'): p.paragraph_format.keep_with_next = True
    return p

def table(rows, widths, first_header=True):
    t = doc.add_table(rows=0, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el = OxmlElement('w:'+side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
    props.append(borders)
    for idx, row in enumerate(rows):
        cells = t.add_row().cells
        trpr = t.rows[-1]._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if idx == 0 and first_header: trpr.append(OxmlElement('w:tblHeader'))
        for j,(cell,text) in enumerate(zip(cells,row)):
            cell.width = Inches(widths[j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr(); margins = OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),'85'); e.set(qn('w:type'),'dxa'); margins.append(e)
            pr.append(margins)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), RED if idx==0 and first_header else ('F4F4F4' if idx%2 else 'FFFFFF')); pr.append(shade)
            p=inline(cell.paragraphs[0],text); p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=1
            for run in p.runs:
                run.font.size=Pt(10.5)
                if idx==0 and first_header: run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

para('INFO XXXX: Digital Experience Engineering','Title')
para('Fall 2026 | Northeastern University College of Engineering','Subtitle')
doc.add_heading('Course Information',1)
opening_rows=[['Course Title','Digital Experience Engineering'],['Course Number','INFO 7375'],['Term and Year','Fall 2026'],['Credit Hours','4 SH'],['CRN','16549'],['Course Format','Online']]
table(opening_rows,[2.05,4.95])
doc.add_heading('Instructor Information',1)
for text in ['Professor Nik Bear Brown | ni.brown@neu.edu | Zoom by Appointment','Professor Nina Harris | n.harris@northeastern.edu | Zoom by Appointment']:
    para(text,'List Bullet')
doc.add_heading('Teaching Assistant Information',1)
for text in ['Full Name: TBD','Email: TBD','Office Hours: TBD']: para(text,'List Bullet')
doc.add_heading('Course Prerequisites',1)
for text in ['Working proficiency in at least one programming language (Python recommended).','Basic familiarity with software development, data, or information systems concepts.','Prior exposure to contemporary development and prototyping tools is helpful but not required.']: para(text,'List Bullet')
doc.add_page_break()

source=(BASE/'digital-experience-engineering-syllabus-proposal.md').read_text()
body=source[source.index('## Course description'):]
body=re.sub(r'## Prerequisites\n.*?(?=## Learning outcomes)', '', body, flags=re.S)
body=body.replace('## Assessment and grading — proposed replacement','## Assessment and Grading')
body=body.replace('The following course-level weights replace the implementation-led categories in the supplied revision for purposes of this proposal. They are not an additional grading layer or a change to an already published course policy.','Course assessment follows six submissions that develop one project from problem framing through design, agentic implementation, evaluation, and delivery.')
body=body.replace("The supplied syllabus's target response time is 48 hours during the academic week.", 'The target response time is 48 hours during the academic week.')
policy_note='The academic-policy text and grade scale below are retained from the supplied proposal. They have not been independently verified as current university policy and require the ordinary institutional syllabus review before student distribution.'
body=body.replace(policy_note+'\n\n','')
body=body.replace('## Faculty review appendix: rationale and teaching resources','## Faculty Review Appendix\n\n'+policy_note.replace('below','in the preceding section'))
body=body.replace('Confirm the permanent course number and title, CRN and modality,', 'Confirm permanent-course approval under the title Digital Experience Engineering, the course number, CRN and modality,')
lines=body.splitlines(); i=0; section=''
while i<len(lines):
    line=lines[i].strip(); i+=1
    if not line or line=='---': continue
    if line.startswith('#'):
        depth=len(line)-len(line.lstrip('#')); title=line.lstrip('# ').strip()
        section=title
        heading=doc.add_heading(title, min(depth-1,3))
        if title in ['Faculty Review Appendix','Administrative provisions']: heading.paragraph_format.page_break_before=True
    elif line.startswith('|'):
        block=[line]
        while i<len(lines) and lines[i].strip().startswith('|'): block.append(lines[i].strip()); i+=1
        rows=[[c.strip() for c in row.strip('|').split('|')] for row in block]
        rows=[row for row in rows if not all(re.fullmatch(r'[-: ]+',c) for c in row)]
        if rows[0][0]=='Week':
            for row in rows[1:]:
                doc.add_heading(f'Week {row[0]}  {row[1]}',2)
                p=para(row[2]); p.paragraph_format.keep_with_next=True
                para('**Evidence or milestone:** '+row[3])
        elif section in ['Playlist alignment','Planned Claude for Educational AI companion playlist','Relationship to the companion courses']:
            for row in rows[1:]:
                doc.add_heading(row[0],3)
                for j,value in enumerate(row[1:]):
                    p=para(value)
                    if j<len(row[1:])-1: p.paragraph_format.keep_with_next=True
        else:
            widths=[1.65,.85,4.5] if rows[0][0]=='Submission' else [2.8,1.3,2.9]
            table(rows,widths)
    elif line.startswith('- '): para(line[2:],'List Bullet')
    elif re.match(r'^\d+\. ',line):
        # Explicit numbers preserve separate list sequences on Google Docs import.
        para(line)
    else:
        if line.startswith('**Religious Observance.'):
            doc.add_heading('University policies and resources',1).paragraph_format.page_break_before=True
        para(line)

footer=s.footer.paragraphs[0]
footer.alignment=2
r=footer.add_run(); r.font.size=Pt(9)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); r._r.addnext(field)
doc.core_properties.title='INFO XXXX: Digital Experience Engineering'
doc.core_properties.subject='Fall 2026 syllabus with faculty review appendix'
doc.core_properties.author='Nik Bear Brown and Nina Harris'
doc.save(BASE/'digital-experience-engineering-fall-2026-unsanitized.docx')
assert [[c.text for c in row.cells] for row in doc.tables[0].rows]==opening_rows
print('Created syllabus; opening course table verified exactly.')
